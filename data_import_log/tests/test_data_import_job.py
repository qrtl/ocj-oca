# Copyright 2026 Quartile (https://www.quartile.co)
# License LGPL-3.0 or later (https://www.gnu.org/licenses/lgpl).

from unittest.mock import patch

from odoo.addons.queue_job.exception import RetryableJobError

from .common import DataImportCase

MODEL = "odoo.addons.data_import_log.models.data_import_log.DataImportLog"


class TestDataImportJob(DataImportCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        # Run the unit and finalizer jobs inline instead of enqueuing them.
        cls.env = cls.env(context=dict(cls.env.context, queue_job__no_delay=True))

    def test_all_units_imported(self):
        log = self._create_log()
        log._start_processing(2)
        with patch(f"{MODEL}._import_unit", return_value=[]):
            log._run_unit("D-001", [])
            self.assertEqual(log.state, "processing")
            log._run_unit("D-002", [])
        self.assertEqual(log.state, "done")
        self.assertTrue(log.date_done)
        self.assertFalse(log.error_ids)

    def test_rejected_unit_makes_the_file_partial(self):
        log = self._create_log()
        log._start_processing(2)
        errors = [{"reference": "D-002", "error_message": "no matching order"}]
        with patch(f"{MODEL}._import_unit", return_value=[]):
            log._run_unit("D-001", [])
        with patch(f"{MODEL}._import_unit", return_value=errors):
            log._run_unit("D-002", [])
        self.assertEqual(log.state, "partial")
        self.assertEqual(log.unit_failed, 1)
        self.assertEqual(log.error_ids.error_message, "no matching order")

    def test_every_unit_rejected_makes_the_file_error(self):
        log = self._create_log()
        log._start_processing(1)
        errors = [{"reference": "D-001", "error_message": "unknown product"}]
        with patch(f"{MODEL}._import_unit", return_value=errors):
            log._run_unit("D-001", [])
        self.assertEqual(log.state, "error")

    def test_transient_failure_settles_nothing(self):
        log = self._create_log()
        log._start_processing(1)
        with patch(f"{MODEL}._import_unit", side_effect=RetryableJobError("locked")):
            with self.assertRaises(RetryableJobError):
                log._run_unit("D-001", [])
        self.assertEqual(log.unit_settled, 0)
        self.assertEqual(log.state, "processing")

    def test_failed_job_settles_its_unit(self):
        log = self._create_log()
        log._start_processing(2)
        with patch(f"{MODEL}._import_unit", return_value=[]):
            log._run_unit("D-001", [])
        job = self.env["queue.job"].create(
            {
                "uuid": "test-failed-unit",
                "model_name": "data.import.log",
                "method_name": "_run_unit",
                "records": log,
                "args": (),
                "kwargs": {"unit_key": "D-002", "rows": []},
                "state": "started",
                "exc_info": "Traceback: boom",
            }
        )
        job.write({"state": "failed"})
        self.assertEqual(log.state, "partial")
        self.assertEqual(log.unit_failed, 1)
        self.assertEqual(log.error_ids.reference, "D-002")
        self.assertIn("boom", log.error_ids.error_message)

    def test_finalize_runs_once(self):
        log = self._create_log()
        log._start_processing(1)
        with patch(f"{MODEL}._import_unit", return_value=[]):
            log._run_unit("D-001", [])
        date_done = log.date_done
        log._finalize()
        self.assertEqual(log.date_done, date_done)

    def test_parse_file_makes_one_unit_per_row(self):
        log = self._create_log(content=b"a,b\n1,2\n3,4\n")
        with patch(f"{MODEL}._import_unit", return_value=[]) as import_unit:
            log._parse_file()
        self.assertEqual(log.unit_total, 2)
        self.assertEqual(log.state, "done")
        self.assertEqual(import_unit.call_count, 2)

    def test_parse_file_groups_rows(self):
        log = self._create_log(content=b"key,qty\nD-1,2\nD-1,3\nD-2,4\n")

        def group(self, fieldnames, rows):
            units = {}
            for row in rows:
                units.setdefault(row["key"], []).append(row)
            return units

        with (
            patch(f"{MODEL}._group_rows", group),
            patch(f"{MODEL}._import_unit", return_value=[]),
        ):
            log._parse_file()
        self.assertEqual(log.unit_total, 2)
        self.assertEqual(log.state, "done")

    def test_empty_file_is_done_not_error(self):
        log = self._create_log(content=b"a,b\n")
        log._parse_file()
        self.assertEqual(log.unit_total, 0)
        self.assertEqual(log.state, "done")
        self.assertFalse(log.error_ids)

    def test_unreadable_file_is_flagged_and_closed(self):
        # CP932 content read as UTF-8: the file itself cannot be read.
        log = self._create_log(content="伝票,数量\nD-1,2\n".encode("cp932"))
        log._parse_file()
        self.assertTrue(log.file_error)
        self.assertEqual(log.state, "error")
        self.assertEqual(log.unit_total, 0)
        self.assertTrue(log.error_ids)

    def test_rejected_units_are_written_back(self):
        content = b"key,qty\nD-1,2\nD-2,3\n"
        log = self._create_log(content=content)
        errors = [{"error_message": "no matching order"}]

        def only_second(self, unit_key, rows):
            return errors if unit_key == "2" else []

        with patch(f"{MODEL}._import_unit", only_second):
            log._parse_file()
        self.assertEqual(log.state, "partial")
        fieldnames, rows = log._rejected_rows()
        self.assertEqual(fieldnames, ["key", "qty"])
        self.assertEqual(rows, [{"key": "D-2", "qty": "3"}])
        name, written = log._rejected_file()
        self.assertEqual(written, b"key,qty\r\nD-2,3\r\n")

    def test_nothing_is_written_back_when_all_units_import(self):
        log = self._create_log(content=b"key,qty\nD-1,2\n")
        with patch(f"{MODEL}._import_unit", return_value=[]):
            log._parse_file()
        self.assertIsNone(log._rejected_file())

    def test_rejected_units_are_written_back_as_xlsx(self):
        import io as _io

        import openpyxl

        workbook = openpyxl.Workbook()
        workbook.active.append(["key", "qty"])
        workbook.active.append(["D-1", 2])
        stream = _io.BytesIO()
        workbook.save(stream)
        log = self._create_log(
            content=stream.getvalue(), file_name="feed.xlsx", file_format="xlsx"
        )
        with patch(f"{MODEL}._import_unit", return_value=[{"error_message": "no"}]):
            log._parse_file()
        name, written = log._rejected_file()
        back = openpyxl.load_workbook(_io.BytesIO(written))
        self.assertEqual(
            [list(r) for r in back.active.iter_rows(values_only=True)],
            [["key", "qty"], ["D-1", "2"]],
        )
