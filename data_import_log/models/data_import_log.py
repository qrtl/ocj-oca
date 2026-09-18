# Copyright 2020-2026 Quartile (https://www.quartile.co)
# License LGPL-3.0 or later (https://www.gnu.org/licenses/lgpl).

import base64
import csv
import datetime
import hashlib
import io

from odoo import api, fields, models
from odoo.exceptions import UserError

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None


class DataImportLog(models.Model):
    _name = "data.import.log"
    _inherit = ["mail.thread", "mail.activity.mixin"]
    _description = "Data Import Log"
    _rec_name = "file_name"
    _order = "id DESC"

    company_id = fields.Many2one(
        "res.company", "Company", default=lambda self: self.env.company
    )
    attachment_id = fields.Many2one(
        "ir.attachment", string="File", readonly=True, ondelete="restrict"
    )
    file_name = fields.Char(related="attachment_id.name", store=True)
    file_data = fields.Binary(related="attachment_id.datas", string="File Content")
    content_hash = fields.Char(
        readonly=True,
        help="SHA-256 of the file content, used to recognize a file already taken in.",
    )
    file_format = fields.Selection(
        [("csv", "CSV"), ("xlsx", "Excel")], required=True, default="csv"
    )
    encoding = fields.Char(
        default="utf-8",
        help="Character encoding of the source file. Ignored for Excel files.",
    )
    model_id = fields.Many2one("ir.model", string="Model")
    model_name = fields.Char(related="model_id.model", string="Model Name")
    state = fields.Selection(
        [
            ("pending", "Pending"),
            ("processing", "Processing"),
            ("done", "Done"),
            ("partial", "Partially Imported"),
            ("error", "Error"),
        ],
        string="Status",
        required=True,
        default="pending",
        readonly=True,
        tracking=True,
    )
    unit_total = fields.Integer(
        "Units", readonly=True, help="Number of units (records or groups) to import."
    )
    unit_settled = fields.Integer(
        "Settled", readonly=True, help="Units that reached a final outcome."
    )
    unit_failed = fields.Integer("Failed", readonly=True)
    date_start = fields.Datetime(
        "Started On", readonly=True, default=fields.Datetime.now
    )
    date_done = fields.Datetime("Finished On", readonly=True)
    error_ids = fields.One2many("data.import.error", "log_id", string="Log Lines")

    _file_uniq = models.Constraint(
        "UNIQUE (file_name, content_hash)",
        "This file has already been taken in.",
    )

    @api.model
    def _content_hash(self, content):
        """Return the hash identifying a file content."""
        return hashlib.sha256(content).hexdigest()

    def _start_processing(self, unit_total):
        """Record how many units the file was split into, and await settlement."""
        self.ensure_one()
        if self.state != "pending":
            raise UserError(
                self.env._(
                    "Import log %(name)s is in state %(state)s and cannot be started.",
                    name=self.display_name,
                    state=self.state,
                )
            )
        self.write({"state": "processing", "unit_total": unit_total})

    def _settle_unit(self, failed=False):
        """Record one unit as having reached a final outcome.

        Return whether this was the last unit of the file, which makes the
        caller responsible for finalizing the log.

        The counters are moved by a single statement so that the row lock
        serializes the workers settling units of the same file concurrently;
        a read-modify-write through the ORM would lose increments. It runs in
        the caller's transaction on purpose: a unit that is rolled back after
        an error must not count as settled either.
        """
        self.ensure_one()
        # The counters are read back from the table, so any pending ORM value
        # (unit_total, most of all) has to be in it first.
        self.flush_recordset(["unit_total", "unit_settled", "unit_failed"])
        self.env.cr.execute(
            """
            UPDATE data_import_log
               SET unit_settled = COALESCE(unit_settled, 0) + 1,
                   unit_failed = COALESCE(unit_failed, 0) + %s
             WHERE id = %s
            RETURNING unit_settled, COALESCE(unit_total, 0)
            """,
            (1 if failed else 0, self.id),
        )
        settled, total = self.env.cr.fetchone()
        self.invalidate_recordset(["unit_settled", "unit_failed"])
        return settled >= total

    @api.model
    def _normalize_cell(self, value):
        """Return a spreadsheet cell as the string a CSV would have held.

        Excel hands back typed values, so the same feed read as xlsx and as CSV
        would otherwise reach the handlers as different Python types.
        """
        if value is None:
            return ""
        if isinstance(value, bool):
            return str(value)
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        if isinstance(value, (datetime.datetime, datetime.date)):
            return value.isoformat()
        return str(value).strip()

    def _read_rows_csv(self, content):
        # utf-8-sig also covers plain utf-8, and keeps a BOM out of the first header.
        encoding = self.encoding or "utf-8"
        if encoding.lower().replace("_", "-") == "utf-8":
            encoding = "utf-8-sig"
        try:
            text = content.decode(encoding)
        except (UnicodeDecodeError, LookupError) as err:
            raise UserError(
                self.env._(
                    "File %(name)s could not be read as %(encoding)s: %(error)s",
                    name=self.file_name,
                    encoding=self.encoding,
                    error=err,
                )
            ) from err
        reader = csv.DictReader(io.StringIO(text))
        fieldnames = [name.strip() for name in reader.fieldnames or []]
        rows = [
            {
                name: self._normalize_cell(value)
                for name, value in zip(fieldnames, row.values(), strict=False)
            }
            for row in reader
        ]
        return fieldnames, rows

    def _read_rows_xlsx(self, content):
        if openpyxl is None:  # pragma: no cover
            raise UserError(
                self.env._(
                    "The openpyxl library is required to read %s.", self.file_name
                )
            )
        workbook = openpyxl.load_workbook(
            io.BytesIO(content), read_only=True, data_only=True
        )
        sheet = workbook[workbook.sheetnames[0]]
        rows_iter = sheet.iter_rows(values_only=True)
        header = next(rows_iter, None)
        if header is None:
            return [], []
        fieldnames = [self._normalize_cell(cell) for cell in header]
        rows = []
        for row in rows_iter:
            values = [self._normalize_cell(cell) for cell in row]
            if not any(values):  # trailing rows Excel keeps around
                continue
            rows.append(dict(zip(fieldnames, values, strict=False)))
        workbook.close()
        return fieldnames, rows

    def _read_rows(self):
        """Return the file content as ``(fieldnames, rows)``.

        ``rows`` are dicts keyed by the header, with every value a string, so a
        handler reads a CSV feed and an Excel feed the same way.
        """
        self.ensure_one()
        content = base64.b64decode(self.attachment_id.datas or b"")
        return getattr(self, f"_read_rows_{self.file_format}")(content)

    def _import_unit(self, unit_key, rows):
        """Import one unit of the file, and return the errors it was rejected for.

        To be implemented by the module that knows the data: return an empty
        list when the unit is imported, a list of values for
        ``data.import.error`` when it is rejected as a whole, and raise
        ``RetryableJobError`` when the failure is transient and the unit should
        be attempted again.
        """
        raise NotImplementedError

    def _enqueue_unit(self, unit_key, rows):
        """Schedule one unit of the file for import."""
        self.ensure_one()
        description = self.env._(
            "Import %(file)s: %(unit)s", file=self.file_name, unit=unit_key
        )
        self.with_delay(description=description)._run_unit(unit_key=unit_key, rows=rows)

    def _run_unit(self, unit_key, rows):
        """Import one unit and account for it, whatever its outcome."""
        self.ensure_one()
        errors = self._import_unit(unit_key, rows)
        if errors:
            self.env["data.import.error"].create(
                [dict(error, log_id=self.id) for error in errors]
            )
        if self._settle_unit(failed=bool(errors)):
            self._enqueue_finalizer()

    def _settle_failed_unit(self, unit_key, reason):
        """Account for a unit whose job failed outside of its own transaction."""
        self.ensure_one()
        self.env["data.import.error"].create(
            {
                "log_id": self.id,
                "reference": unit_key or "",
                "error_message": reason,
            }
        )
        if self._settle_unit(failed=True):
            self._enqueue_finalizer()

    def _enqueue_finalizer(self):
        self.ensure_one()
        description = self.env._("Finalize import of %(file)s", file=self.file_name)
        self.with_delay(description=description)._finalize()

    def _finalize_file(self):
        """Hook for disposing of the source file, extended where it is stored."""

    def _notify_outcome(self):
        """Hook for reporting an import that did not fully succeed."""

    def _finalize(self):
        """Close the log once every unit has settled.

        Guarded on the state, so that settling a unit twice, or the sweeper
        racing a last unit, cannot finalize the same file twice.
        """
        self.ensure_one()
        if self.state != "processing":
            return
        if self.unit_failed >= self.unit_total:
            state = "error"
        elif self.unit_failed:
            state = "partial"
        else:
            state = "done"
        self.write({"state": state, "date_done": fields.Datetime.now()})
        self._finalize_file()
        if state != "done":
            self._notify_outcome()
